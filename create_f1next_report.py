from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


# ============================================================
# CREATE DOCUMENT
# ============================================================

doc = Document()

# Page margins
section = doc.sections[0]
section.top_margin = Inches(0.7)
section.bottom_margin = Inches(0.7)
section.left_margin = Inches(0.8)
section.right_margin = Inches(0.8)


# ============================================================
# HELPER FUNCTIONS
# ============================================================

def set_cell_text(cell, text, bold=False):
    cell.text = str(text)

    for paragraph in cell.paragraphs:
        for run in paragraph.runs:
            run.bold = bold
            run.font.size = Pt(9)

    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def add_table(headers, rows):
    table = doc.add_table(
        rows=1,
        cols=len(headers)
    )

    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"

    # Header
    for i, header in enumerate(headers):
        set_cell_text(
            table.rows[0].cells[i],
            header,
            bold=True
        )

    # Rows
    for row in rows:
        cells = table.add_row().cells

        for i, value in enumerate(row):
            set_cell_text(cells[i], value)

    doc.add_paragraph()

    return table


def add_heading(text, level=1):
    doc.add_heading(text, level=level)


def add_bullet(text, level=0):
    paragraph = doc.add_paragraph(
        style="List Bullet" if level == 0 else "List Bullet 2"
    )
    paragraph.add_run(text)


def add_number(text):
    paragraph = doc.add_paragraph(
        style="List Number"
    )
    paragraph.add_run(text)


# ============================================================
# TITLE
# ============================================================

title = doc.add_paragraph()
title.alignment = WD_ALIGN_PARAGRAPH.CENTER

run = title.add_run(
    "F1Next\n"
    "Machine Learning-Based Formula 1 Race and Championship Prediction"
)

run.bold = True
run.font.size = Pt(22)

subtitle = doc.add_paragraph()
subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER

run = subtitle.add_run(
    "Project Progress and Review 1 Documentation"
)

run.italic = True
run.font.size = Pt(13)

doc.add_paragraph()


# ============================================================
# 1. PROJECT OVERVIEW
# ============================================================

add_heading("1. Project Overview", 1)

doc.add_paragraph(
    "F1Next is a machine learning project focused on analyzing historical "
    "Formula 1 race and qualifying data to predict driver race performance."
)

doc.add_paragraph(
    "The project currently contains two prediction tasks:"
)

add_bullet(
    "Regression: Predict the numerical finishing position of a driver."
)

add_bullet(
    "Classification: Predict whether a driver will finish on the podium."
)

doc.add_paragraph(
    "The project follows a chronological forecasting approach so that "
    "future-season information is not used during model training."
)


# ============================================================
# 2. OBJECTIVES
# ============================================================

add_heading("2. Project Objectives", 1)

objectives = [
    "Collect and integrate historical Formula 1 race and qualifying data.",
    "Perform exploratory data analysis on the dataset.",
    "Clean and preprocess the data.",
    "Engineer historical driver and constructor performance features.",
    "Build multiple regression models for finishing-position prediction.",
    "Build multiple classification models for podium prediction.",
    "Compare machine learning algorithms using appropriate evaluation metrics.",
    "Perform cross-validation and hyperparameter tuning.",
    "Evaluate the models on an unseen future season.",
    "Prepare the system for future 2026 Formula 1 prediction."
]

for item in objectives:
    add_number(item)


# ============================================================
# 3. DATASET
# ============================================================

add_heading("3. Dataset", 1)

doc.add_paragraph(
    "Historical Formula 1 race and qualifying datasets were collected "
    "for the seasons 2010 through 2025."
)

add_table(
    ["Dataset", "Seasons", "Approximate Records"],
    [
        ["Race Results", "2010–2025", "6,911"],
        ["Qualifying Results", "2010–2025", "6,891"]
    ]
)

doc.add_paragraph(
    "Race and qualifying data were merged using the following keys:"
)

add_bullet("season")
add_bullet("round")
add_bullet("driver_id")

doc.add_paragraph(
    "The merged dataset contains 6,911 race-driver records."
)


# ============================================================
# 4. DATA INTEGRATION
# ============================================================

add_heading("4. Data Integration", 1)

doc.add_paragraph(
    "Race-result and qualifying datasets were combined into a single "
    "dataset using season, race round, and driver ID."
)

doc.add_paragraph(
    "Duplicate checks were performed before and after merging. "
    "No duplicate season-round-driver records were found."
)

add_table(
    ["Check", "Result"],
    [
        ["Race duplicate records", "0"],
        ["Qualifying duplicate records", "0"],
        ["Merged dataset", "6,911 rows"]
    ]
)


# ============================================================
# 5. EDA
# ============================================================

add_heading("5. Exploratory Data Analysis (EDA)", 1)

doc.add_paragraph(
    "Exploratory Data Analysis was performed to understand the structure, "
    "distribution, and relationships within the dataset before model "
    "development."
)

add_heading("5.1 Dataset Audit", 2)

add_table(
    ["Property", "Result"],
    [
        ["Rows", "6,911"],
        ["Columns", "35"],
        ["Duplicate rows", "0"]
    ]
)

add_heading("5.2 Missing Value Analysis", 2)

doc.add_paragraph(
    "The final dataset contains missing values in several raw qualifying "
    "and fastest-lap attributes. These attributes were not included in "
    "the final 14-feature modeling set."
)

add_table(
    ["Column", "Missing Values"],
    [
        ["fastest_lap_rank", "287"],
        ["fastest_lap_speed", "5,614"],
        ["fastest_lap_time", "2,071"],
        ["Q1", "114"],
        ["Q2", "1,905"],
        ["Q3", "3,743"]
    ]
)

doc.add_paragraph(
    "The selected 14 modeling features were subsequently checked and "
    "contained zero missing values after preprocessing."
)

add_heading("5.3 Target Distribution", 2)

doc.add_paragraph(
    "The regression target is finish position. The distribution is "
    "relatively consistent across the main finishing positions, with "
    "fewer observations occurring at the later positions."
)

doc.add_paragraph(
    "The classification target is podium:"
)

add_bullet("1 = Podium finish (1st, 2nd, or 3rd)")
add_bullet("0 = No podium finish")

doc.add_paragraph(
    "The podium target is imbalanced, with substantially more non-podium "
    "observations than podium observations. Therefore, classification "
    "evaluation uses Accuracy together with Weighted F1 and ROC-AUC."
)

add_heading("5.4 Feature Distribution Observations", 2)

add_bullet(
    "Qualifying position and grid position are distributed across the available positions."
)

add_bullet(
    "Historical finishing and qualifying averages are concentrated around middle positions."
)

add_bullet(
    "Season and constructor points, wins, and podium counts are right-skewed."
)

add_bullet(
    "Career races before the current race are also right-skewed."
)

add_heading("5.5 Correlation Analysis", 2)

doc.add_paragraph(
    "The correlation heatmap showed meaningful relationships between "
    "qualifying performance, historical performance, accumulated points, "
    "and finishing position."
)

add_table(
    ["Feature", "Correlation with Finish Position"],
    [
        ["qualifying_position", "0.63"],
        ["grid", "0.60"],
        ["avg_qualifying_last_5", "0.61"],
        ["avg_finish_last_5", "0.58"],
        ["avg_finish_all_previous", "0.54"],
        ["avg_qualifying_all_previous", "0.54"],
        ["season_points_so_far", "-0.46"],
        ["constructor_points_so_far", "-0.46"]
    ]
)

doc.add_paragraph(
    "Several engineered features also showed strong correlations with "
    "each other, indicating some degree of redundancy among historical "
    "performance variables."
)

add_heading("5.6 Feature-Target Relationships", 2)

add_bullet(
    "Qualifying position showed a positive relationship with finishing position."
)

add_bullet(
    "Recent qualifying performance also showed a positive relationship "
    "with race finishing position."
)

doc.add_paragraph(
    "These relationships indicate that qualifying and previous performance "
    "contain useful information for predicting race outcomes, although "
    "substantial variation remains due to other race-related factors."
)


# ============================================================
# 6. PREPROCESSING
# ============================================================

add_heading("6. Data Preprocessing", 1)

steps = [
    "Missing-value analysis",
    "Duplicate checking",
    "Data type conversion",
    "Invalid-value checking",
    "Chronological sorting",
    "Historical feature engineering",
    "Driver performance feature engineering",
    "Constructor performance feature engineering",
    "Career experience feature engineering",
    "Outlier analysis using the IQR method",
    "Classification target creation",
    "Feature selection",
    "Feature scaling"
]

for step in steps:
    add_bullet(step)

add_heading("6.1 Outlier Analysis", 2)

doc.add_paragraph(
    "The IQR method was used to identify potential outliers. "
    "Extreme values were not automatically removed because some extreme "
    "values represent legitimate Formula 1 performance conditions."
)

add_heading("6.2 Feature Scaling", 2)

doc.add_paragraph(
    "StandardScaler was fitted only on the training dataset and then "
    "applied to validation and test datasets."
)

doc.add_paragraph(
    "This prevents information from future datasets from influencing "
    "the preprocessing stage."
)


# ============================================================
# 7. FEATURE ENGINEERING
# ============================================================

add_heading("7. Feature Engineering", 1)

doc.add_paragraph(
    "Historical features were engineered using previous race information "
    "to avoid using the current race result as an input."
)

feature_columns = [
    "qualifying_position",
    "grid",
    "previous_finish",
    "avg_finish_all_previous",
    "avg_qualifying_all_previous",
    "avg_finish_last_5",
    "avg_qualifying_last_5",
    "season_points_so_far",
    "season_wins_so_far",
    "season_podiums_so_far",
    "constructor_points_so_far",
    "constructor_wins_so_far",
    "constructor_podiums_so_far",
    "career_races_before"
]

for feature in feature_columns:
    add_bullet(feature)


# ============================================================
# 8. DATA SPLIT
# ============================================================

add_heading("8. Chronological Data Split", 1)

doc.add_paragraph(
    "A chronological split was used because the project is designed "
    "for future-season forecasting."
)

add_table(
    ["Dataset", "Seasons", "Records"],
    [
        ["Training", "2010–2023", "5,953"],
        ["Validation", "2024", "479"],
        ["Testing", "2025", "479"],
        ["Future Prediction", "2026", "Future"]
    ]
)

doc.add_paragraph(
    "The 2025 test dataset was kept separate from model selection and "
    "hyperparameter tuning."
)


# ============================================================
# 9. REGRESSION
# ============================================================

add_heading("9. Regression Modeling", 1)

doc.add_paragraph(
    "The regression task predicts the numerical race finishing position "
    "of each driver."
)

add_heading("9.1 Regression Models", 2)

regression_models = [
    "Linear Regression",
    "Ridge Regression",
    "Lasso Regression",
    "ElasticNet Regression",
    "Polynomial Regression",
    "Decision Tree Regressor",
    "Random Forest Regressor",
    "Gradient Boosting Regressor",
    "Support Vector Regression (SVR)",
    "K-Nearest Neighbors (KNN) Regressor"
]

for model in regression_models:
    add_bullet(model)

add_heading("9.2 Regression Evaluation Metrics", 2)

add_bullet("R²")
add_bullet("RMSE")
add_bullet("MAE")


# ============================================================
# 10. REGRESSION VALIDATION RESULTS
# ============================================================

add_heading("10. Regression Validation Results — 2024", 1)

add_table(
    ["Model", "R²", "RMSE", "MAE"],
    [
        ["Lasso Regression", "0.5670", "3.7874", "2.9472"],
        ["ElasticNet Regression", "0.5668", "3.7886", "2.9439"],
        ["Ridge Regression", "0.5658", "3.7928", "2.9368"],
        ["Linear Regression", "0.5657", "3.7931", "2.9368"],
        ["Polynomial Regression", "0.5538", "3.8449", "3.0167"],
        ["Gradient Boosting", "0.5317", "3.9389", "3.0990"],
        ["Random Forest", "0.5095", "4.0311", "3.1660"],
        ["KNN", "0.4362", "4.3218", "3.3712"],
        ["SVR", "0.4312", "4.3410", "3.1816"],
        ["Decision Tree", "-0.2615", "6.4649", "4.8434"]
    ]
)

doc.add_paragraph(
    "The initial model comparison was performed on the 2024 validation "
    "dataset. Lasso and ElasticNet were selected for further cross-validation "
    "based on the highest validation R² values."
)


# ============================================================
# 11. CROSS VALIDATION
# ============================================================

add_heading("11. Regression Cross-Validation", 1)

doc.add_paragraph(
    "Five-fold chronological cross-validation was performed on the "
    "2010–2023 training data for the two selected models."
)

add_table(
    ["Model", "Fold 1", "Fold 2", "Fold 3", "Fold 4", "Fold 5", "Mean R²"],
    [
        [
            "Lasso",
            "0.3846",
            "0.4432",
            "0.4062",
            "0.4007",
            "0.4265",
            "0.4122"
        ],
        [
            "ElasticNet",
            "0.3884",
            "0.4438",
            "0.4078",
            "0.4006",
            "0.4259",
            "0.4133"
        ]
    ]
)


# ============================================================
# 12. HYPERPARAMETER TUNING
# ============================================================

add_heading("12. Hyperparameter Tuning", 1)

doc.add_paragraph(
    "GridSearchCV with chronological TimeSeriesSplit was used to tune "
    "the two selected regression models."
)

add_table(
    ["Model", "Best Parameters", "Best CV R²"],
    [
        ["Lasso", "alpha = 0.05", "0.4150"],
        [
            "ElasticNet",
            "alpha = 0.05, l1_ratio = 0.9",
            "0.4151"
        ]
    ]
)

doc.add_paragraph(
    "The tuned models showed a small improvement in cross-validation "
    "performance compared with their initial settings."
)


# ============================================================
# 13. TUNED VALIDATION RESULTS
# ============================================================

add_heading("13. Tuned Regression Results — 2024 Validation", 1)

add_table(
    ["Model", "R²", "RMSE", "MAE"],
    [
        ["Tuned Lasso", "0.5676", "3.7849", "2.9587"],
        ["Tuned ElasticNet", "0.5675", "3.7856", "2.9582"]
    ]
)

doc.add_paragraph(
    "Tuned Lasso was selected for the final regression test evaluation "
    "based on its slightly higher 2024 validation R²."
)


# ============================================================
# 14. REGRESSION VISUALIZATIONS
# ============================================================

add_heading("14. Regression Visualizations", 1)

add_bullet(
    "Actual vs Predicted plot was generated for the Tuned Lasso model."
)

add_bullet(
    "Residual plot was generated to analyze prediction errors."
)

add_bullet(
    "Random Forest feature importance was generated to identify influential features."
)

doc.add_paragraph(
    "The Actual vs Predicted plot showed that predictions generally "
    "follow the actual finishing-position trend, while also showing "
    "prediction concentration toward middle finishing positions."
)

doc.add_paragraph(
    "The residual plot showed residuals around zero but also displayed "
    "a noticeable pattern and several larger residuals, indicating "
    "systematic prediction errors for certain finishing positions."
)


# ============================================================
# 15. FEATURE IMPORTANCE
# ============================================================

add_heading("15. Random Forest Feature Importance", 1)

add_table(
    ["Feature", "Importance"],
    [
        ["qualifying_position", "0.3567"],
        ["avg_qualifying_last_5", "0.1248"],
        ["avg_finish_all_previous", "0.0775"],
        ["avg_finish_last_5", "0.0708"],
        ["avg_qualifying_all_previous", "0.0693"],
        ["career_races_before", "0.0641"],
        ["constructor_points_so_far", "0.0608"],
        ["previous_finish", "0.0469"],
        ["season_points_so_far", "0.0418"],
        ["grid", "0.0396"],
        ["constructor_podiums_so_far", "0.0180"],
        ["season_podiums_so_far", "0.0130"],
        ["constructor_wins_so_far", "0.0092"],
        ["season_wins_so_far", "0.0073"]
    ]
)

doc.add_paragraph(
    "Qualifying position had the highest Random Forest feature importance, "
    "followed by recent qualifying performance and previous finishing performance."
)


# ============================================================
# 16. FINAL REGRESSION TEST
# ============================================================

add_heading("16. Final Regression Test — 2025", 1)

doc.add_paragraph(
    "The 2025 dataset was kept untouched during training, validation, "
    "model comparison, and hyperparameter tuning."
)

add_table(
    ["Metric", "Tuned Lasso — 2025 Test"],
    [
        ["R²", "0.4518"],
        ["RMSE", "4.2618"],
        ["MAE", "3.2128"]
    ]
)

doc.add_paragraph(
    "The MAE of 3.2128 indicates that the predicted finishing position "
    "differs from the actual finishing position by approximately 3.21 "
    "positions on average on the 2025 test data."
)


# ============================================================
# 17. CLASSIFICATION
# ============================================================

add_heading("17. Classification Modeling", 1)

doc.add_paragraph(
    "The classification task predicts whether a driver finishes on "
    "the podium."
)

doc.add_paragraph(
    "Target definition:"
)

add_bullet("1 = Podium (1st, 2nd, or 3rd)")
add_bullet("0 = No Podium")

add_heading("17.1 Classification Models", 2)

classification_models = [
    "Logistic Regression",
    "K-Nearest Neighbors (KNN)",
    "Gaussian Naive Bayes",
    "Decision Tree",
    "Support Vector Classifier (SVC)"
]

for model in classification_models:
    add_bullet(model)

add_heading("17.2 Classification Metrics", 2)

add_bullet("Accuracy")
add_bullet("Weighted F1-score")
add_bullet("ROC-AUC")
add_bullet("Confusion Matrix")


# ============================================================
# 18. CLASSIFICATION VALIDATION
# ============================================================

add_heading("18. Classification Validation Results — 2024", 1)

add_table(
    ["Model", "Accuracy", "Weighted F1", "ROC-AUC"],
    [
        ["Logistic Regression", "0.8831", "0.8786", "0.9271"],
        ["SVC", "0.8831", "0.8810", "0.9028"],
        ["KNN", "0.8643", "0.8682", "0.8869"],
        ["Gaussian Naive Bayes", "0.8058", "0.8286", "0.8826"],
        ["Decision Tree", "0.8539", "0.8577", "0.7425"]
    ]
)

doc.add_paragraph(
    "The classification models were compared using accuracy, weighted "
    "F1-score, and ROC-AUC on the 2024 validation dataset."
)


# ============================================================
# 19. CONFUSION MATRIX SUMMARY
# ============================================================

add_heading("19. Classification Confusion Matrices", 1)

doc.add_paragraph(
    "Confusion matrices were generated for all five classification models "
    "using the 2024 validation dataset."
)

add_table(
    ["Model", "TN", "FP", "FN", "TP"],
    [
        ["Logistic Regression", "385", "22", "34", "38"],
        ["KNN", "369", "38", "27", "45"],
        ["Gaussian Naive Bayes", "325", "82", "11", "61"],
        ["Decision Tree", "367", "40", "30", "42"],
        ["SVC", "382", "25", "31", "41"]
    ]
)


# ============================================================
# 20. FINAL CLASSIFICATION TEST
# ============================================================

add_heading("20. Final Classification Test — 2025", 1)

doc.add_paragraph(
    "The 2025 dataset was evaluated after the models had been trained "
    "using the 2010–2023 training period."
)

add_table(
    ["Model", "Accuracy", "Weighted F1", "ROC-AUC"],
    [
        ["Logistic Regression", "0.9269", "0.9263", "0.9411"],
        ["KNN", "0.8643", "0.8694", "0.8920"],
        ["Gaussian Naive Bayes", "0.8601", "0.8727", "0.9183"],
        ["Decision Tree", "0.8810", "0.8820", "0.7756"],
        ["SVC", "0.9186", "0.9193", "0.9308"]
    ]
)

doc.add_paragraph(
    "Logistic Regression achieved 0.9269 accuracy, 0.9263 weighted "
    "F1-score, and 0.9411 ROC-AUC on the unseen 2025 test dataset."
)


# ============================================================
# 21. MODELS SELECTED
# ============================================================

add_heading("21. Models Selected for Further Development", 1)

add_heading("21.1 Regression", 2)

doc.add_paragraph(
    "Lasso and ElasticNet were selected for cross-validation and "
    "hyperparameter tuning because they produced the two highest "
    "validation R² values among the initial ten regression models."
)

doc.add_paragraph(
    "After tuning, Tuned Lasso was used for the final 2025 regression "
    "test evaluation because it achieved a slightly higher 2024 "
    "validation R² of 0.5676 compared with Tuned ElasticNet at 0.5675."
)

add_heading("21.2 Classification", 2)

doc.add_paragraph(
    "All five classification models were evaluated on the unseen 2025 "
    "test season. Logistic Regression produced the strongest overall "
    "2025 test metrics in the current experiment, including an "
    "accuracy of 0.9269 and ROC-AUC of 0.9411."
)

doc.add_paragraph(
    "These results will be used as the baseline for future development "
    "and should not be interpreted as a guarantee of future-season performance."
)


# ============================================================
# 22. CURRENT WORKFLOW
# ============================================================

add_heading("22. Current F1Next Workflow", 1)

workflow = [
    "Historical F1 data collection",
    "Race and qualifying data integration",
    "EDA",
    "Data preprocessing",
    "Historical feature engineering",
    "Chronological train/validation/test split",
    "Regression modeling",
    "Classification Part A modeling",
    "Model evaluation",
    "Cross-validation",
    "Hyperparameter tuning",
    "2025 final testing",
    "Future 2026 prediction"
]

for i, step in enumerate(workflow, start=1):
    add_number(step)


# ============================================================
# 23. CURRENT STATUS
# ============================================================

add_heading("23. Project Status", 1)

add_table(
    ["Component", "Status"],
    [
        ["Data Collection", "Completed"],
        ["Data Integration", "Completed"],
        ["EDA", "Completed"],
        ["Preprocessing", "Completed"],
        ["Feature Engineering", "Completed"],
        ["Regression — 10 Models", "Completed"],
        ["Regression Evaluation", "Completed"],
        ["Regression Cross-Validation", "Completed"],
        ["Regression Hyperparameter Tuning", "Completed"],
        ["Regression Final Test", "Completed"],
        ["Classification Part A — 5 Models", "Completed"],
        ["Classification Evaluation", "Completed"],
        ["Classification Confusion Matrices", "Completed"],
        ["Classification Final Test", "Completed"],
        ["README", "In Progress"],
        ["Review 2", "Pending"],
        ["2026 Prediction", "Pending"]
    ]
)


# ============================================================
# 24. FUTURE WORK
# ============================================================

add_heading("24. Future Work", 1)

future_work = [
    "Complete Review 2 classification models.",
    "Implement additional ensemble classification algorithms.",
    "Perform clustering using K-Means.",
    "Perform Agglomerative Hierarchical Clustering.",
    "Evaluate clustering using appropriate metrics.",
    "Apply PCA and t-SNE for dimensionality reduction and visualization.",
    "Develop the 2026 prediction pipeline.",
    "Generate driver race predictions for future races.",
    "Develop a final user-facing dashboard/application.",
    "Prepare final project presentation and documentation."
]

for item in future_work:
    add_bullet(item)


# ============================================================
# 25. AI ASSISTANCE
# ============================================================

add_heading("25. AI Assistance", 1)

doc.add_paragraph(
    "AI assistance was used during development for code scaffolding, "
    "debugging support, explanation of machine learning concepts, "
    "notebook organization, and implementation guidance."
)

doc.add_paragraph(
    "The dataset processing, execution of the machine learning models, "
    "evaluation results, and project implementation were performed "
    "and verified during project development."
)


# ============================================================
# 26. CONCLUSION
# ============================================================

add_heading("26. Current Conclusion", 1)

doc.add_paragraph(
    "The current F1Next implementation successfully establishes a "
    "machine learning pipeline for Formula 1 race-performance prediction. "
    "Historical data from 2010–2025 has been integrated and processed, "
    "and historical driver and constructor performance features have "
    "been engineered."
)

doc.add_paragraph(
    "Ten regression algorithms and five classification algorithms were "
    "implemented and evaluated. Cross-validation and hyperparameter "
    "tuning were also performed for the regression task."
)

doc.add_paragraph(
    "The 2025 season was kept as an unseen test dataset. The current "
    "results provide a baseline for continuing the project toward "
    "additional models, clustering, and 2026 future prediction."
)


# ============================================================
# SAVE
# ============================================================

output_file = "F1Next_Progress_Report.docx"

doc.save(output_file)

print("==============================================")
print("F1Next documentation created successfully!")
print("File:", output_file)
print("==============================================")